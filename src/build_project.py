"""Digital Payment Insight — reproducible synthetic-data project.
Run from the repository root: python src/build_project.py
All records/labels are synthetic; this is an educational prototype, not a fraud system.
"""
from pathlib import Path
import json
import warnings

import joblib
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import RandomForestClassifier
from sklearn.impute import SimpleImputer
from sklearn.metrics import accuracy_score, precision_score, recall_score, f1_score, roc_auc_score, confusion_matrix
from sklearn.model_selection import train_test_split
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

warnings.filterwarnings("ignore", category=FutureWarning)
ROOT = Path(__file__).resolve().parents[1]
RNG = np.random.default_rng(49113)
N = 12000
for folder in ("data/raw", "data/processed", "reports/figures", "models", "excel", "presentation"):
    (ROOT / folder).mkdir(parents=True, exist_ok=True)

# Generate reference tables.
n_customers, n_merchants, n_devices = 1800, 400, 3600
customer_ids = np.array([f"C{i:05d}" for i in range(1, n_customers + 1)])
merchant_ids = np.array([f"M{i:04d}" for i in range(1, n_merchants + 1)])
device_ids = np.array([f"D{i:05d}" for i in range(1, n_devices + 1)])
segments = RNG.choice(["Student", "Regular", "Premium", "Business"], n_customers, p=[.22, .48, .22, .08])
regions = RNG.choice(["North", "South", "East", "West", "Central"], n_customers, p=[.20, .27, .15, .23, .15])
customers = pd.DataFrame({
    "customer_id": customer_ids, "customer_segment": segments, "customer_region": regions,
    "tenure_months": np.clip(RNG.gamma(2.5, 18, n_customers).astype(int), 1, 120)
})
categories = np.array(["Grocery", "Dining", "Travel", "Electronics", "Utilities", "Fashion", "Gaming", "Healthcare", "Fuel", "Other"])
merchants = pd.DataFrame({
    "merchant_id": merchant_ids, "merchant_category": RNG.choice(categories, n_merchants),
    "merchant_region": RNG.choice(["North", "South", "East", "West", "Central"], n_merchants)
})
devices = pd.DataFrame({
    "device_id": device_ids, "device_type": RNG.choice(["Android", "iOS", "Web", "Tablet"], n_devices, p=[.46, .28, .20, .06]),
    "new_device_flag": RNG.choice([0, 1], n_devices, p=[.84, .16]),
    "device_history_complete": RNG.choice([1, 0], n_devices, p=[.94, .06])
})

# Create 30 days of transactions with an evening peak and injected anomalies.
hour_probs = np.array([.01,.005,.003,.002,.002,.003,.01,.025,.055,.065,.06,.06,.065,.06,.055,.05,.05,.06,.075,.085,.075,.055,.03,.015])
hour_probs = hour_probs / hour_probs.sum()
hours = RNG.choice(np.arange(24), N, p=hour_probs)
times = (pd.Timestamp("2026-09-01") + pd.to_timedelta(RNG.integers(0, 30, N), unit="D")
         + pd.to_timedelta(hours, unit="h") + pd.to_timedelta(RNG.integers(0, 3600, N), unit="s"))
ci, mi, di = RNG.integers(0, n_customers, N), RNG.integers(0, n_merchants, N), RNG.integers(0, n_devices, N)
segment, category = segments[ci], merchants.merchant_category.to_numpy()[mi]
new_device = devices.new_device_flag.to_numpy()[di]
region_mismatch = (regions[ci] != merchants.merchant_region.to_numpy()[mi])
base = {"Student": 250, "Regular": 650, "Premium": 1500, "Business": 2200}
amount = np.array([base[x] for x in segment]) * RNG.lognormal(0, .72, N)
amount *= np.select([category == "Travel", category == "Electronics", category == "Utilities", category == "Grocery", category == "Gaming"], [2.1, 2.6, .45, .55, .8], default=1.0)
is_night = np.isin(hours, [0, 1, 2, 3, 4])
spike = RNG.random(N) < .035
amount = np.where(spike, amount * RNG.uniform(5, 12, N), amount)
amount = np.where((new_device == 1) & is_night, amount * RNG.uniform(1.4, 2.5, N), amount)
log_amount = np.log1p(amount)
threshold = np.select([segment == "Student", segment == "Regular", segment == "Premium", segment == "Business"], [7.9, 8.6, 9.3, 9.8])
risk = 1.15*(log_amount > threshold) + 1.1*(new_device == 1) + is_night + .7*region_mismatch + 1.5*spike + RNG.normal(0, .5, N)
prob = 1 / (1 + np.exp(-(risk - 2.25)))
labels = (RNG.random(N) < prob).astype(int)
channels = RNG.choice(["Mobile App", "Web", "UPI", "Card", "mobile app", "WEB", "upi"], N, p=[.30,.16,.22,.12,.08,.06,.06])
transactions = pd.DataFrame({
    "transaction_id": [f"T{i:07d}" for i in range(1, N+1)], "customer_id": customer_ids[ci],
    "merchant_id": merchant_ids[mi], "device_id": device_ids[di], "transaction_time": times,
    "amount": np.round(amount, 2), "merchant_category": category, "channel": channels
})
missing_idx = RNG.choice(N, int(N * .025), replace=False)
transactions.loc[missing_idx, "merchant_category"] = np.nan
transactions = pd.concat([transactions, transactions.sample(35, random_state=13)], ignore_index=True)
label_map = dict(zip([f"T{i:07d}" for i in range(1, N+1)], labels))
unique_tx = transactions.drop_duplicates("transaction_id", keep="first")
label_series = unique_tx.transaction_id.map(label_map).fillna(0).astype(int)
alerts = pd.DataFrame({
    "transaction_id": unique_tx.transaction_id,
    "review_status": np.where(label_series.to_numpy() == 1, RNG.choice(["Reviewed", "Pending"], len(unique_tx), p=[.7,.3]), "No review"),
    "reason_category": np.where(label_series.to_numpy() == 1, RNG.choice(["High amount", "New device", "Unusual hour", "Region mismatch", "Multiple signals"], len(unique_tx)), "Routine"),
    "synthetic_review_label": label_series.to_numpy()
})
for name, frame in [("transactions", transactions), ("customers", customers), ("devices", devices), ("merchants", merchants), ("alerts", alerts)]:
    frame.to_csv(ROOT / "data" / "raw" / f"{name}.csv", index=False)

# Clean, standardize and enrich.
clean = transactions.drop_duplicates("transaction_id", keep="first").copy()
raw_rows = len(transactions)
clean["transaction_time"] = pd.to_datetime(clean["transaction_time"], errors="coerce")
clean["channel"] = (clean["channel"].astype(str).str.strip().str.lower()
                    .replace({"mobile app": "mobile_app", "web": "web", "upi": "upi", "card": "card"}))
clean["merchant_category"] = clean["merchant_category"].fillna("Unknown")
cat_map = merchants.set_index("merchant_id")["merchant_category"]
unknown = clean["merchant_category"].eq("Unknown")
clean.loc[unknown, "merchant_category"] = clean.loc[unknown, "merchant_id"].map(cat_map).fillna("Unknown")
clean["amount"] = pd.to_numeric(clean["amount"], errors="coerce")
amount_cap = clean["amount"].quantile(.995)
clean["amount_capped"] = clean["amount"].clip(upper=amount_cap)
clean["hour"] = clean["transaction_time"].dt.hour
clean["day_of_week"] = clean["transaction_time"].dt.day_name()
clean["is_night"] = clean["hour"].isin([0,1,2,3,4]).astype(int)
clean = clean.merge(customers, on="customer_id", how="left")
clean = clean.merge(devices[["device_id","device_type","new_device_flag","device_history_complete"]], on="device_id", how="left")
clean = clean.merge(merchants[["merchant_id","merchant_region"]], on="merchant_id", how="left")
clean["region_mismatch"] = (clean["customer_region"] != clean["merchant_region"]).astype(int)
clean["log_amount"] = np.log1p(clean["amount_capped"].clip(lower=0))
clean["synthetic_review_label"] = clean["transaction_id"].map(label_map).fillna(0).astype(int)
clean = clean.merge(alerts[["transaction_id","review_status","reason_category"]], on="transaction_id", how="left")
clean.to_csv(ROOT / "data/processed/cleaned_transactions.csv", index=False)

segment_summary = clean.groupby("customer_segment").agg(transaction_count=("transaction_id","count"), median_amount=("amount","median"), p95_amount=("amount", lambda x:x.quantile(.95)), review_rate=("synthetic_review_label","mean")).reset_index()
hour_summary = clean.groupby("hour").agg(transaction_count=("transaction_id","count"), review_rate=("synthetic_review_label","mean"), median_amount=("amount","median")).reset_index()
category_summary = clean.groupby("merchant_category").agg(transaction_count=("transaction_id","count"), review_count=("synthetic_review_label","sum"), review_rate=("synthetic_review_label","mean"), median_amount=("amount","median")).reset_index().sort_values("review_rate", ascending=False)
device_summary = clean.groupby("new_device_flag").agg(transaction_count=("transaction_id","count"), review_rate=("synthetic_review_label","mean"), median_amount=("amount","median")).reset_index()
region_summary = clean.groupby("customer_region").agg(transaction_count=("transaction_id","count"), review_count=("synthetic_review_label","sum"), review_rate=("synthetic_review_label","mean")).reset_index()
summaries = {"segment_summary": segment_summary, "hour_summary": hour_summary, "merchant_category_summary": category_summary, "device_summary": device_summary, "region_summary": region_summary}
for name, frame in summaries.items():
    frame.to_csv(ROOT / "data/processed" / f"{name}.csv", index=False)

# Train/evaluate an interpretable baseline pipeline; labels are synthetic proxies only.
features = ["log_amount","hour","is_night","new_device_flag","region_mismatch","tenure_months","customer_segment","merchant_category","channel","device_type","customer_region"]
target = "synthetic_review_label"
X, y = clean[features], clean[target]
numeric = ["log_amount","hour","is_night","new_device_flag","region_mismatch","tenure_months"]
categorical = ["customer_segment","merchant_category","channel","device_type","customer_region"]
preprocess = ColumnTransformer([
    ("numeric", Pipeline([("imputer", SimpleImputer(strategy="median")), ("scale", StandardScaler())]), numeric),
    ("categorical", Pipeline([("imputer", SimpleImputer(strategy="most_frequent")), ("onehot", OneHotEncoder(handle_unknown="ignore"))]), categorical)
])
model = Pipeline([("preprocess", preprocess), ("classifier", RandomForestClassifier(n_estimators=180, min_samples_leaf=3, class_weight="balanced_subsample", random_state=49113, n_jobs=-1))])
X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=.2, random_state=49113, stratify=y)
model.fit(X_train, y_train)
pred = model.predict(X_test)
score = model.predict_proba(X_test)[:,1]
metrics = {
    "dataset_rows_after_cleaning": int(len(clean)), "raw_transaction_rows": int(raw_rows),
    "duplicate_rows_removed": int(raw_rows-len(clean)), "missing_merchant_categories_filled": int(len(missing_idx)),
    "test_rows": int(len(X_test)), "target_positive_rate": float(y.mean()),
    "accuracy": float(accuracy_score(y_test,pred)), "precision": float(precision_score(y_test,pred,zero_division=0)),
    "recall": float(recall_score(y_test,pred,zero_division=0)), "f1_score": float(f1_score(y_test,pred,zero_division=0)),
    "roc_auc": float(roc_auc_score(y_test,score)), "confusion_matrix": confusion_matrix(y_test,pred).tolist(),
    "threshold_note": "Default classifier threshold 0.50; synthetic proxy labels only."
}
(ROOT / "reports/model_metrics.json").write_text(json.dumps(metrics, indent=2))
joblib.dump(model, ROOT / "models/synthetic_review_model.joblib")

# EDA charts.
sns.set_theme(style="whitegrid")
plt.figure(figsize=(9,5)); sns.histplot(clean["amount"], bins=60, log_scale=(False, True)); plt.title("Transaction Amount Distribution"); plt.xlabel("Amount (log-scaled count axis)"); plt.tight_layout(); plt.savefig(ROOT/"reports/figures/amount_distribution.png", dpi=150); plt.close()
plt.figure(figsize=(10,5)); sns.barplot(data=hour_summary, x="hour", y="transaction_count", color="#3973ac"); plt.title("Transaction Volume by Hour"); plt.xlabel("Hour of Day"); plt.ylabel("Transactions"); plt.tight_layout(); plt.savefig(ROOT/"reports/figures/hourly_volume.png", dpi=150); plt.close()
plt.figure(figsize=(7,5)); sns.barplot(data=device_summary, x="new_device_flag", y="review_rate", color="#db8b35"); plt.xticks([0,1],["Known device","New device"]); plt.ylabel("Synthetic review rate"); plt.title("Review Rate: New vs Known Device"); plt.tight_layout(); plt.savefig(ROOT/"reports/figures/device_comparison.png", dpi=150); plt.close()
plt.figure(figsize=(10,5)); sns.barplot(data=category_summary, x="merchant_category", y="review_rate", color="#568b67"); plt.xticks(rotation=35, ha="right"); plt.ylabel("Synthetic review rate"); plt.title("Review Rate by Merchant Category"); plt.tight_layout(); plt.savefig(ROOT/"reports/figures/merchant_category_review_rate.png", dpi=150); plt.close()
plt.figure(figsize=(8,5)); sns.barplot(data=region_summary, x="customer_region", y="review_rate", color="#6c6fb0"); plt.ylabel("Synthetic review rate"); plt.title("Review Rate by Customer Region"); plt.tight_layout(); plt.savefig(ROOT/"reports/figures/regional_trends.png", dpi=150); plt.close()
plt.figure(figsize=(5,4)); sns.heatmap(confusion_matrix(y_test,pred), annot=True, fmt="d", cmap="Blues", cbar=False); plt.xlabel("Predicted"); plt.ylabel("Actual synthetic label"); plt.title("Holdout Confusion Matrix"); plt.tight_layout(); plt.savefig(ROOT/"reports/figures/confusion_matrix.png", dpi=150); plt.close()

peak = hour_summary.loc[hour_summary.transaction_count.idxmax()]
new_rate = float(device_summary.loc[device_summary.new_device_flag == 1, "review_rate"].iloc[0])
known_rate = float(device_summary.loc[device_summary.new_device_flag == 0, "review_rate"].iloc[0])
topcat = category_summary.iloc[0]
findings = {
    "peak_hour": int(peak.hour), "peak_hour_transactions": int(peak.transaction_count),
    "new_device_review_rate": new_rate, "known_device_review_rate": known_rate,
    "highest_category_by_synthetic_review_rate": str(topcat.merchant_category),
    "highest_category_review_rate": float(topcat.review_rate),
    "note": "Synthetic data and synthetic review-priority labels; not evidence of actual fraud."
}
(ROOT / "reports/key_findings.json").write_text(json.dumps(findings, indent=2))

# Excel workbook with summary, analysis tabs, cleaned records and charts.
wb = Workbook()
ws = wb.active
ws.title = "Executive Summary"
navy = "17365D"; blue = "D9EAF7"; gold = "FCE4D6"
ws["A1"] = "DIGITAL PAYMENT INSIGHT — SYNTHETIC DATA"; ws["A1"].font = Font(bold=True, size=16, color="FFFFFF"); ws["A1"].fill = PatternFill("solid", fgColor=navy)
ws.merge_cells("A1:D1")
summary_rows = [
    ("Metric","Value"),("Raw transactions",raw_rows),("Clean transactions",len(clean)),("Duplicate records removed",raw_rows-len(clean)),
    ("Missing merchant categories imputed",len(missing_idx)),("Accuracy",metrics["accuracy"]),("Precision",metrics["precision"]),
    ("Recall",metrics["recall"]),("F1-score",metrics["f1_score"]),("ROC-AUC",metrics["roc_auc"]),
    ("Peak hour",f"{int(peak.hour):02d}:00"),("Transactions at peak hour",int(peak.transaction_count)),
    ("New-device synthetic review rate",new_rate),("Known-device synthetic review rate",known_rate),
]
for row in summary_rows: ws.append(row)
for cell in ws[2]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor=navy)
for row in ws.iter_rows(min_row=3):
    row[0].font=Font(bold=True)
    if row[0].value and "rate" in str(row[0].value).lower(): row[1].number_format="0.00%"
ws["A17"]="Important: all records and labels are synthetic. Metrics do not measure real fraud detection."
ws.merge_cells("A17:D17"); ws["A17"].font=Font(italic=True,color="9C0006")
for sheet_name, frame in [("Cleaned Data",clean),("By Segment",segment_summary),("By Hour",hour_summary),("By Merchant",category_summary),("By Device",device_summary),("By Region",region_summary)]:
    sh=wb.create_sheet(sheet_name)
    for row in [list(frame.columns)] + frame.astype(object).where(pd.notna(frame), "").values.tolist(): sh.append(row)
    sh.freeze_panes="A2"; sh.auto_filter.ref=sh.dimensions
    for cell in sh[1]:
        cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor=navy)
    for col in sh.columns:
        max_len=min(max(len(str(c.value or "")) for c in list(col)[:100])+2,28)
        sh.column_dimensions[col[0].column_letter].width=max_len
    for col_name in frame.columns:
        if "rate" in str(col_name).lower():
            col_idx=list(frame.columns).index(col_name)+1
            for rr in range(2, sh.max_row+1): sh.cell(rr,col_idx).number_format="0.00%"
# chart sheets
for sheet_name, source_sheet, xcol, ycol, title in [
    ("Chart Hourly","By Hour",1,2,"Transaction Volume by Hour"),
    ("Chart Merchant","By Merchant",1,4,"Synthetic Review Rate by Category"),
    ("Chart Region","By Region",1,4,"Synthetic Review Rate by Region")
]:
    sh=wb.create_sheet(sheet_name); sh["A1"]=title; sh["A1"].font=Font(bold=True,size=14)
    src=wb[source_sheet]
    chart=BarChart(); chart.title=title; chart.y_axis.title="Count" if "Volume" in title else "Review rate"
    data=Reference(src,min_col=ycol,min_row=1,max_row=src.max_row)
    cats_ref=Reference(src,min_col=xcol,min_row=2,max_row=src.max_row)
    chart.add_data(data,titles_from_data=True); chart.set_categories(cats_ref); chart.height=8; chart.width=16; sh.add_chart(chart,"A3")
wb.save(ROOT/"excel/digital_payment_analysis.xlsx")

# Pitch deck.
prs=Presentation()
def slide(title, bullets):
    s=prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text=title
    tf=s.placeholders[1].text_frame; tf.clear()
    for i, item in enumerate(bullets):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=item; p.level=0; p.font.size=Pt(22)
slide("Digital Payment Insight", ["Use Case 13 | DS Sprint Assessment 2", "Synthetic educational dataset for transaction monitoring", "Goal: understand patterns that deserve operational review"])
slide("Business Question & Data", [f"{raw_rows:,} raw rows → {len(clean):,} cleaned transactions", "Five tables: transactions, customers, devices, merchants, alerts", "Injected quality issues: duplicates, missing categories and inconsistent channel labels"])
slide("Cleaning & Feature Engineering", ["Removed duplicate transaction IDs and imputed missing merchant categories", "Normalized channels, parsed timestamps and joined reference tables", "Features: log amount, hour, night flag, new device, region mismatch and customer tenure"])
slide("EDA Findings", [f"Highest-volume hour: {int(peak.hour):02d}:00 ({int(peak.transaction_count):,} transactions)", f"New-device review rate: {new_rate:.1%}; known-device: {known_rate:.1%}", f"Highest-category synthetic review rate: {topcat.merchant_category} ({topcat.review_rate:.1%})"])
slide("Machine Learning", ["Algorithm: Random Forest Classifier", f"Accuracy {metrics['accuracy']:.3f} | Precision {metrics['precision']:.3f} | Recall {metrics['recall']:.3f}", f"F1 {metrics['f1_score']:.3f} | ROC-AUC {metrics['roc_auc']:.3f}", "Scores predict synthetic proxy labels only; results are modest"])
slide("Monitoring Priorities", ["Review combinations of new-device activity, unusually high amount, night-time activity and region mismatch", "Use category/hour/region charts to guide analyst attention", "Avoid automatic payment blocking based on this prototype"])
slide("Recommendations", ["Validate thresholds with confirmed, historical outcomes", "Measure false positives and false negatives with domain experts", "Keep human review and governance in the decision loop"])
slide("Conclusion & Limitations", ["End-to-end synthetic data, EDA, ML, Excel and reporting workflow", "Educational prototype only — not validated fraud detection", "Replace synthetic records/labels and retest before operational use"])
prs.save(ROOT/"presentation/digital_payment_insight_pitch.pptx")

print(json.dumps({"output_directory": str(ROOT), "raw_rows": raw_rows, "clean_rows": len(clean), "metrics": metrics, "findings": findings}, indent=2))
